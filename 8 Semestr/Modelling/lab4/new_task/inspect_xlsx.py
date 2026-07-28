import os
from pathlib import Path

from openpyxl import load_workbook


def dump_wb(path: Path, max_rows: int = 40, max_cols: int = 12) -> None:
    print(f"\n=== {path.name} ===")
    wb = load_workbook(path, data_only=False)
    print("sheets:", wb.sheetnames)
    for s in wb.sheetnames:
        ws = wb[s]
        # list some non-empty cells (formulas/instructions)
        nonempty = []
        for row in ws.iter_rows(values_only=False):
            for cell in row:
                v = cell.value
                if v not in (None, ""):
                    if isinstance(v, str):
                        vv = " ".join(v.split())
                    else:
                        vv = v
                    nonempty.append((cell.coordinate, vv))

        print(f"--- sheet: {s} --- nonempty={len(nonempty)}")
        for coord, vv in nonempty[:60]:
            print(coord, vv)
        if len(nonempty) > 60:
            print("... (more non-empty cells)")

        for r in range(1, max_rows + 1):
            row = []
            for c in range(1, max_cols + 1):
                v = ws.cell(row=r, column=c).value
                if isinstance(v, str):
                    v = " ".join(v.split())
                row.append(v)
            if any(x not in (None, "") for x in row):
                print(r, row)


def main() -> None:
    base = Path(__file__).resolve().parents[1]
    new_task = base / "new_task"
    files = sorted([new_task / f for f in os.listdir(new_task) if f.lower().endswith(".xlsx")])
    print("xlsx files:", [f.name for f in files])
    for f in files:
        dump_wb(f)


if __name__ == "__main__":
    main()

