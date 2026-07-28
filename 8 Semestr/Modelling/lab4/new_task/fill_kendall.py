from __future__ import annotations

import math
import os
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def kendall_from_xy(x: list[float], y: list[float]) -> dict[str, float]:
    n = len(x)
    n0 = n * (n - 1) // 2

    nc = 0
    nd = 0
    for i in range(n):
        for j in range(i + 1, n):
            dx = x[i] - x[j]
            dy = y[i] - y[j]
            prod = dx * dy
            if prod > 0:
                nc += 1
            elif prod < 0:
                nd += 1

    cx = Counter(x)
    cy = Counter(y)
    tx = sum(v * (v - 1) // 2 for v in cx.values())
    ty = sum(v * (v - 1) // 2 for v in cy.values())

    tau_a = (nc - nd) / n0 if n0 else 0.0
    denom = math.sqrt((n0 - tx) * (n0 - ty)) if n0 else 0.0
    tau_b = (nc - nd) / denom if denom else 0.0

    z = tau_b * math.sqrt(9 * n * (n - 1) / (2 * (2 * n + 5))) if n > 1 else 0.0
    # two-sided p-value via erfc (normal tail), but we will NOT write it to xlsx
    p = math.erfc(abs(z) / math.sqrt(2.0))

    return {
        "n": float(n),
        "n0": float(n0),
        "nc": float(nc),
        "nd": float(nd),
        "tx": float(tx),
        "ty": float(ty),
        "tau_a": float(tau_a),
        "tau_b": float(tau_b),
        "z": float(z),
        "p": float(p),
    }


def apply_kendall(path: Path, alpha: float = 0.05) -> None:
    wb = load_workbook(path, data_only=False)
    if len(wb.sheetnames) < 2:
        raise RuntimeError("Workbook must have at least 2 sheets.")

    s1, s2 = wb.sheetnames[0], wb.sheetnames[1]
    ws1 = wb[s1]

    # Recreate sheet2 (compute-only), do NOT touch sheet1 (Spearman)
    old_ws2 = wb[s2]
    wb.remove(old_ws2)
    ws2 = wb.create_sheet(title=s2, index=1)

    # Copy the original data columns to sheet2 (values only)
    ws2["A1"] = ws1["A1"].value or "Рост в см"
    ws2["B1"] = ws1["B1"].value or "Вес в кг"
    raw: list[tuple[float, float]] = []
    for r in range(2, 19):
        xv = ws1[f"A{r}"].value
        yv = ws1[f"B{r}"].value
        if xv is None or yv is None:
            continue
        raw.append((float(xv), float(yv)))

    # Variant B: sort by X (Рост), then by Y (Вес) for stability
    raw_sorted = sorted(raw, key=lambda t: (t[0], t[1]))
    for i, (xv, yv) in enumerate(raw_sorted, start=2):
        ws2[f"A{i}"] = xv
        ws2[f"B{i}"] = yv
    # Clear remaining rows if any (template has 17 rows)
    for r in range(2 + len(raw_sorted), 19):
        ws2[f"A{r}"] = None
        ws2[f"B{r}"] = None

    # Add C/D/E columns with formulas (as in the earlier table)
    ws2["C1"] = "C (соглас.)"
    ws2["D1"] = "D (несогл.)"
    ws2["E1"] = "T (связ.)"

    last_row = 1 + len(raw_sorted)  # data rows: 2..last_row
    for r in range(2, last_row + 1):
        r_next = r + 1
        if r_next > last_row:
            ws2[f"C{r}"] = 0
            ws2[f"D{r}"] = 0
            ws2[f"E{r}"] = 0
            continue
        # No locale-specific separators used here (SUMPRODUCT has one argument)
        ws2[f"C{r}"] = (
            f"=SUMPRODUCT(--((($A${r_next}:$A${last_row}-A{r})*($B${r_next}:$B${last_row}-B{r}))>0))"
        )
        ws2[f"D{r}"] = (
            f"=SUMPRODUCT(--((($A${r_next}:$A${last_row}-A{r})*($B${r_next}:$B${last_row}-B{r}))<0))"
        )
        ws2[f"E{r}"] = (
            f"=SUMPRODUCT(--((($A${r_next}:$A${last_row}-A{r})*($B${r_next}:$B${last_row}-B{r}))=0))"
        )

    x_vals: list[float] = []
    y_vals: list[float] = []
    for r in range(2, last_row + 1):
        xv = ws2[f"A{r}"].value
        yv = ws2[f"B{r}"].value
        if xv is None or yv is None:
            continue
        x_vals.append(float(xv))
        y_vals.append(float(yv))

    st = kendall_from_xy(x_vals, y_vals)

    # Summary block (values; p-value left blank by request)
    ws2["I1"] = "Критерий Кендалла (tau)"
    ws2["I3"] = "n="
    ws2["J3"] = int(st["n"])

    ws2["I4"] = "N0=n(n-1)/2"
    ws2["J4"] = int(st["n0"])

    ws2["I5"] = "Nc (соглас.)"
    ws2["J5"] = int(st["nc"])

    ws2["I6"] = "Nd (несогл.)"
    ws2["J6"] = int(st["nd"])

    ws2["I7"] = "T_x (связи по X)"
    ws2["J7"] = int(st["tx"])

    ws2["I8"] = "T_y (связи по Y)"
    ws2["J8"] = int(st["ty"])

    ws2["I9"] = "tau_a"
    ws2["J9"] = st["tau_a"]

    ws2["I10"] = "tau_b (с поправкой на связи)"
    ws2["J10"] = st["tau_b"]

    ws2["I12"] = "z (H0: tau=0)"
    ws2["J12"] = st["z"]

    ws2["I13"] = "p-value (двуст.)"
    ws2["J13"] = ""  # do not write p-value; user will paste formula manually

    ws2["I14"] = "α"
    ws2["J14"] = alpha

    ws2["I15"] = "Решение"
    ws2["J15"] = "отвергаем H0 (связь есть)" if st["p"] < alpha else "нет оснований отвергнуть H0"

    # widths (cosmetic)
    for col, w in [("A", 10), ("B", 10), ("C", 13), ("D", 13), ("E", 10), ("I", 26), ("J", 24)]:
        ws2.column_dimensions[col].width = w

    wb.save(path)


def pick_target_xlsx(new_task_dir: Path) -> Path:
    xlsx_files = sorted([new_task_dir / f for f in os.listdir(new_task_dir) if f.lower().endswith(".xlsx")])
    if not xlsx_files:
        raise RuntimeError("No .xlsx files found.")

    for f in xlsx_files:
        if "Домрачева" in f.name or "РК" in f.name:
            return f
    return xlsx_files[0]


def main() -> None:
    base = Path(__file__).resolve().parents[1]
    new_task = base / "new_task"
    target = pick_target_xlsx(new_task)
    apply_kendall(target, alpha=0.05)
    print("Updated:", target)


if __name__ == "__main__":
    main()

